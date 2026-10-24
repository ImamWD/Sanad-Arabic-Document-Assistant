#!/usr/bin/env python3
"""Render the seed corpus (C1b).

    make render-seed            (or: docker compose run --rm render python3 seed/render.py)

Runs in the `render` tools image (tools/render/Dockerfile): Chromium, LibreOffice, poppler,
qpdf, img2pdf and python3 with yaml + python-docx. It writes, deterministically:

- seed/documents/<file>  one file per manifest entry, from seed/sources/<doc_id>.md
    pdf   Markdown -> HTML -> Chromium print. One fixed-size A4 box per source page, so a page
          marker is always a page break and a source page can never spill onto the next one.
    docx  python-docx, an explicit page break before the first block of every page after the
          first, RTL paragraphs and tables, Arabic set as the complex-script font.
    txt   the body verbatim, a form feed between pages, encoded as the manifest says.
- seed/bad/              ingestion failure fixtures (see seed/bad/README.md)
- seed/previews/         PNGs of three rendered pages: an Arabic PDF, an English PDF, and a DOCX
                         after LibreOffice conversion

Then it checks every rendered page count against `page_count`, converting DOCX and TXT with
LibreOffice the way ingestion does (scripts/check_eval.py repeats that check and also checks
the text of every page).
"""

# ruff: noqa: S603, S607, T201  -- a CLI tool: it prints progress and runs local renderers.
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import html
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SEED = ROOT / "seed"
DOCS_DIR = SEED / "documents"
BAD_DIR = SEED / "bad"
PREVIEW_DIR = SEED / "previews"

# Every timestamp written into a rendered file. Fixed, so re-rendering is byte-identical.
FIXED_DATE = dt.datetime(2026, 10, 24, tzinfo=dt.UTC)
PDF_DATE = b"D:20261024000000Z"

# Pages rendered as PNG previews: (manifest file, page). One Arabic PDF, one English PDF and
# one DOCX after LibreOffice conversion.
PREVIEWS = [
    ("annual-report-2020.pdf", 11),
    ("annual-report-2023-en.pdf", 9),
    ("budget-2018.docx", 3),
]

MARKER = re.compile(r"^<!-- page (\d+) -->[ \t]*$", re.M)
FRONT_MATTER_FIELDS = (
    "doc_id",
    "title_ar",
    "title_en",
    "category",
    "year",
    "programme",
    "language",
    "format",
    "page_count",
    "replaces",
    "encoding",
)

LIBREOFFICE = [
    "soffice",
    "--headless",
    "--norestore",
    "-env:UserInstallation=file:///tmp/lo-profile",
]


class RenderError(Exception):
    pass


# --------------------------------------------------------------------------------------------
# Source parsing
# --------------------------------------------------------------------------------------------


@dataclass
class Source:
    entry: dict[str, object]
    pages: list[str]

    @property
    def doc_id(self) -> str:
        return str(self.entry["doc_id"])

    @property
    def language(self) -> str:
        return str(self.entry["language"])


def load_manifest() -> dict[str, object]:
    return json.loads((SEED / "manifest.json").read_text(encoding="utf-8"))


def load_source(entry: dict[str, object]) -> Source:
    """Read seed/sources/<doc_id>.md, check it against the manifest, split it into pages."""
    path = ROOT / str(entry["source"])
    text = path.read_text(encoding="utf-8")
    if not text.startswith("---\n"):
        raise RenderError(f"{path}: no front matter")
    end = text.index("\n---\n", 4)
    front = yaml.safe_load(text[4:end])
    body = text[end + 5 :]
    for key in FRONT_MATTER_FIELDS:
        if front.get(key) != entry.get(key):
            raise RenderError(
                f"{path}: front matter {key}={front.get(key)!r} != manifest {entry.get(key)!r}"
            )
    numbers = [int(n) for n in MARKER.findall(body)]
    count = int(entry["page_count"])  # type: ignore[arg-type]
    if numbers != list(range(1, count + 1)):
        raise RenderError(f"{path}: page markers {numbers} != 1..{count}")
    parts = MARKER.split(body)
    if parts[0].strip():
        raise RenderError(f"{path}: text before <!-- page 1 -->")
    # split() with one group yields [pre, "1", page1, "2", page2, ...]
    pages = [parts[i] for i in range(2, len(parts), 2)]
    return Source(entry, pages)


# A tiny parser for the Markdown subset in seed/README.md §2: headings, paragraphs, bullet and
# numbered lists, **bold**, blockquotes and pipe tables. One parser feeds both the HTML and the
# DOCX writer, so the two formats can't disagree about what a page contains.


@dataclass
class Block:
    kind: str  # h1 h2 h3 p quote ul ol table
    text: str = ""
    items: list[str] = field(default_factory=list)
    start: int = 1
    header: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


BULLET = re.compile(r"^- (.*)$")
NUMBERED = re.compile(r"^(\d+)\. (.*)$")
HEADING = re.compile(r"^(#{1,3}) (.*)$")


def split_row(line: str) -> list[str]:
    line = line.strip()
    if not (line.startswith("|") and line.endswith("|")):
        raise RenderError(f"malformed table row: {line!r}")
    return [cell.strip() for cell in line[1:-1].split("|")]


def parse_blocks(page: str) -> list[Block]:
    blocks: list[Block] = []
    lines = page.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if m := HEADING.match(line):
            blocks.append(Block(f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
        elif line.startswith("|"):
            rows: list[str] = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            if len(rows) < 2 or not re.fullmatch(r"\|(\s*:?-+:?\s*\|)+", rows[1].strip()):
                raise RenderError(f"table without a header separator: {rows[0]!r}")
            header = split_row(rows[0])
            body = [split_row(r) for r in rows[2:]]
            if any(len(r) != len(header) for r in body):
                raise RenderError(f"ragged table under {rows[0]!r}")
            blocks.append(Block("table", header=header, rows=body))
        elif line.startswith(">"):
            quoted: list[str] = []
            while i < len(lines) and lines[i].startswith(">"):
                quoted.append(lines[i][1:].strip())
                i += 1
            blocks.append(Block("quote", " ".join(q for q in quoted if q)))
        elif BULLET.match(line) or NUMBERED.match(line):
            kind = "ul" if BULLET.match(line) else "ol"
            pattern = BULLET if kind == "ul" else NUMBERED
            block = Block(kind)
            if kind == "ol":
                block.start = int(NUMBERED.match(line).group(1))  # type: ignore[union-attr]
            while i < len(lines) and (m := pattern.match(lines[i].rstrip())):
                block.items.append(m.group(m.lastindex or 1).strip())
                i += 1
            blocks.append(block)
        else:
            # Paragraphs are one line each in the sources; a following list interrupts one.
            blocks.append(Block("p", line.strip()))
            i += 1
    return blocks


def inline_runs(text: str) -> list[tuple[str, bool]]:
    """Split `**bold**` markup into (text, bold) runs."""
    parts = text.split("**")
    if len(parts) % 2 == 0:
        raise RenderError(f"unbalanced ** in {text!r}")
    return [(part, i % 2 == 1) for i, part in enumerate(parts) if part]


# --------------------------------------------------------------------------------------------
# PDF: HTML + Chromium
# --------------------------------------------------------------------------------------------

# Arabic text is set in Amiri. Not Noto Naskh Arabic: it draws dotted letters as a base glyph
# plus dot glyphs, so Chromium wraps nearly every Arabic letter in an /ActualText span, and
# pdfium (pypdfium2) then extracts many words letter-reversed. Amiri maps letters to single
# glyphs. Its optional ligatures (تظ, مص, نق, ...) are switched off in CSS: pdfium drops a
# letter from each. Required ones (lam-alef) stay. Never KacstOne: its lam-meem ligature
# extracts as the wrong letters (M0 finding).
FONTS = {
    "ar": '"Amiri", "Noto Serif", serif',
    "mixed": '"Amiri", "Noto Serif", serif',
    "en": '"Noto Serif", "Amiri", serif',
}
BASE_PT = {"ar": 14.0, "mixed": 13.5, "en": 11.5}
LINE_HEIGHT = {"ar": 1.6, "mixed": 1.6, "en": 1.5}

CSS = """
@page { size: A4; margin: 0; }
html { print-color-adjust: exact; -webkit-print-color-adjust: exact; }
body {
  margin: 0; font-family: %(font)s; color: #1d2327; line-height: %(lh)s;
  font-variant-ligatures: no-common-ligatures no-discretionary-ligatures
    no-historical-ligatures no-contextual;
}
section.page {
  box-sizing: border-box; inline-size: 210mm; block-size: 297mm; overflow: hidden;
  padding-block: 20mm 18mm; padding-inline: 20mm; break-after: page;
}
section.page:last-child { break-after: auto; }
h1, h2, h3 { color: #1f4e5f; line-height: 1.35; margin-block: 0 0.55em; font-weight: 700; }
h1 { font-size: 1.75em; margin-block-end: 0.6em; }
h2 { font-size: 1.3em; margin-block-start: 0.9em; }
h3 { font-size: 1.1em; margin-block-start: 0.8em; }
section.page > :first-child { margin-block-start: 0; }
p { margin-block: 0 0.65em; }
ul, ol { margin-block: 0 0.65em; padding-inline-start: 1.6em; }
li { margin-block-end: 0.25em; }
blockquote {
  margin-block: 0 0.9em; margin-inline: 0; padding-block: 0.45em; padding-inline: 0.9em;
  border-inline-start: 1mm solid #b08d57; background: #f6f1e7; font-size: 0.9em;
}
table {
  border-collapse: collapse; inline-size: 100%%; margin-block: 0.2em 0.9em; font-size: 0.92em;
}
th, td {
  border: 0.75pt solid #7d8b91; padding-block: 0.25em; padding-inline: 0.55em;
  text-align: start; vertical-align: top; line-height: 1.45;
}
th { background: #e3ecef; font-weight: 700; }
"""

# Shrinks the type of a page until it fits its A4 box, in 0.25pt steps. Most pages never
# shrink (the sources are written to about half of the page budget). data-fit records the size
# used and data-overflow marks a page that still doesn't fit; render.py fails on that.
FIT_SCRIPT = """
for (const page of document.querySelectorAll("section.page")) {
  let size = parseFloat(page.dataset.base);
  page.style.fontSize = size + "pt";
  while (page.scrollHeight > page.clientHeight && size > 8) {
    size -= 0.25;
    page.style.fontSize = size + "pt";
  }
  page.dataset.fit = String(size);
  if (page.scrollHeight > page.clientHeight) page.dataset.overflow = "1";
}
"""


def html_runs(text: str) -> str:
    return "".join(
        f"<strong>{html.escape(t)}</strong>" if bold else html.escape(t)
        for t, bold in inline_runs(text)
    )


# Mixed documents put an Arabic and an English clause in one block ("عنوان — Title", "البيان /
# Item", an Arabic sentence followed by its English translation). Each clause is isolated in a
# <bdi>, so it takes its own direction and its trailing punctuation stays with it.
CLAUSE_SPLIT = re.compile(r"( — | / |(?<=[.؛:!?])\s+)")


def html_inline(text: str, isolate: bool = False) -> str:
    if not isolate:
        return html_runs(text)
    parts = CLAUSE_SPLIT.split(text)
    if len(parts) == 1:
        return html_runs(text)
    # Bold never spans a separator in the sources; check rather than assume.
    if any(p.count("**") % 2 for p in parts[::2]):
        return html_runs(text)
    return "".join(
        html.escape(p) if i % 2 else f"<bdi>{html_runs(p)}</bdi>" for i, p in enumerate(parts)
    )


def first_strong_dir(text: str) -> str:
    """What dir="auto" computes: the direction of the first strong character."""
    for ch in text:
        bidi = unicodedata.bidirectional(ch)
        if bidi == "L":
            return "ltr"
        if bidi in ("R", "AL"):
            return "rtl"
    return "rtl"


def html_page(blocks: list[Block], mixed: bool) -> str:
    """Mixed documents give each block the direction of its own first strong character. It
    is set explicitly, not as dir="auto": auto skips text inside <bdi>, and in a block whose
    every clause is isolated it would find nothing and fall back to LTR."""

    def d(text: str) -> str:
        return f' dir="{first_strong_dir(text)}"' if mixed else ""

    def inline(text: str) -> str:
        return html_inline(text, isolate=mixed)

    out: list[str] = []
    for b in blocks:
        if b.kind in ("h1", "h2", "h3", "p"):
            out.append(f"<{b.kind}{d(b.text)}>{inline(b.text)}</{b.kind}>")
        elif b.kind == "quote":
            out.append(f"<blockquote{d(b.text)}>{inline(b.text)}</blockquote>")
        elif b.kind in ("ul", "ol"):
            start = f' start="{b.start}"' if b.kind == "ol" and b.start != 1 else ""
            items = "".join(f"<li{d(t)}>{inline(t)}</li>" for t in b.items)
            out.append(f"<{b.kind}{start}>{items}</{b.kind}>")
        elif b.kind == "table":
            head = "".join(f"<th{d(c)}>{inline(c)}</th>" for c in b.header)
            body = "".join(
                "<tr>" + "".join(f"<td{d(c)}>{inline(c)}</td>" for c in row) + "</tr>"
                for row in b.rows
            )
            out.append(f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
    return "\n".join(out)


def build_html(src: Source) -> str:
    lang = src.language
    html_lang, html_dir = ("en", "ltr") if lang == "en" else ("ar", "rtl")
    title = src.entry["title_en"] if lang == "en" else src.entry["title_ar"]
    css = CSS % {"font": FONTS[lang], "lh": LINE_HEIGHT[lang]}
    pages = "\n".join(
        f'<section class="page" data-page="{n}" data-base="{BASE_PT[lang]}">\n'
        f"{html_page(parse_blocks(text), mixed=lang == 'mixed')}\n</section>"
        for n, text in enumerate(src.pages, 1)
    )
    return (
        f'<!doctype html>\n<html lang="{html_lang}" dir="{html_dir}">\n<head>\n'
        f'<meta charset="utf-8">\n<title>{html.escape(str(title))}</title>\n'
        f"<style>{css}</style>\n</head>\n<body>\n{pages}\n"
        f"<script>{FIT_SCRIPT}</script>\n</body>\n</html>\n"
    )


def chromium(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            "chromium",
            "--headless",
            "--no-sandbox",
            "--disable-gpu",
            "--disable-dev-shm-usage",
            "--no-first-run",
            "--user-data-dir=/tmp/chromium-profile",
            *args,
        ],
        check=True,
        capture_output=True,
        text=True,
    )


def fit_report(html_path: Path) -> dict[int, float]:
    """Font size each page ended up with (only pages that had to shrink). Fails on overflow."""
    dom = chromium("--dump-dom", f"file://{html_path}").stdout
    shrunk: dict[int, float] = {}
    for m in re.finditer(r"<section ([^>]*)>", dom):
        attrs = dict(re.findall(r'data-(\w+)="([^"]*)"', m.group(1)))
        page = int(attrs["page"])
        if "overflow" in attrs:
            raise RenderError(f"{html_path.name}: page {page} does not fit on one A4 page")
        if float(attrs["fit"]) < float(attrs["base"]):
            shrunk[page] = float(attrs["fit"])
    return shrunk


def deterministic_pdf(raw: Path, out: Path) -> None:
    """Pin the timestamps Chromium writes, drop its XMP packet and give the file a stable ID."""
    data = raw.read_bytes()
    # Same-length replacement keeps the cross-reference offsets valid for qpdf to read.
    data = re.sub(
        rb"/(CreationDate|ModDate) \((D:[^)]*)\)",
        lambda m: b"/%s (%s)" % (m.group(1), PDF_DATE.ljust(len(m.group(2)))[: len(m.group(2))]),
        data,
    )
    raw.write_bytes(data)
    subprocess.run(
        ["qpdf", "--deterministic-id", "--remove-metadata", str(raw), str(out)],
        check=True,
    )


def render_pdf(src: Source, work: Path, out: Path) -> dict[int, float]:
    html_path = work / f"{src.doc_id}.html"
    html_path.write_text(build_html(src), encoding="utf-8")
    raw = work / f"{src.doc_id}.raw.pdf"
    chromium("--no-pdf-header-footer", f"--print-to-pdf={raw}", f"file://{html_path}")
    shrunk = fit_report(html_path)
    deterministic_pdf(raw, out)
    return shrunk


# --------------------------------------------------------------------------------------------
# DOCX: python-docx
# --------------------------------------------------------------------------------------------

DOCX_FONT_CS = "Noto Naskh Arabic"  # complex script (Arabic)
DOCX_FONT_LATIN = "Noto Serif"
DOCX_PT = {"body": 12.0, "table": 10.5, "notice": 10.5, "h1": 20.0, "h2": 15.0, "h3": 13.0}
HEADING_COLOR = RGBColor(0x1F, 0x4E, 0x5F)


def _set(parent: OxmlElement, tag: str, **attrs: str) -> OxmlElement:
    """Get or create child `tag` of `parent` and set its w: attributes."""
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    for key, value in attrs.items():
        el.set(qn(f"w:{key}"), value)
    return el


def set_rpr(rpr: OxmlElement, size: float | None = None, bold: bool | None = None) -> None:
    """Arabic in Noto Naskh Arabic, Latin in Noto Serif, complex-script size and bold, RTL."""
    fonts = _set(rpr, "w:rFonts")
    for key in list(fonts.attrib):  # drops theme fonts (asciiTheme, cstheme, ...)
        del fonts.attrib[key]
    for slot in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{slot}"), DOCX_FONT_LATIN)
    fonts.set(qn("w:cs"), DOCX_FONT_CS)
    if bold is not None:
        for tag in ("w:b", "w:bCs"):
            _set(rpr, tag, val="1" if bold else "0")
    if size is not None:
        for tag in ("w:sz", "w:szCs"):
            _set(rpr, tag, val=str(round(size * 2)))
    _set(rpr, "w:rtl", val="1")


# Word rejects properties in the wrong child order (ECMA-376 sequences). Elements added with
# _set() are appended, so every pPr/rPr/tblPr/sectPr is put back in schema order at the end.
CHILD_ORDER = {
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss "
    "imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs "
    "highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish",
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr "
    "suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct "
    "topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind "
    "contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
    "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize "
    "tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection "
    "tcFitText vAlign hideMark",
    "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc "
    "pgBorders lnNumType pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi "
    "rtlGutter docGrid printerSettings sectPrChange",
}
CHILD_ORDER_INDEX = {k: {t: i for i, t in enumerate(v.split())} for k, v in CHILD_ORDER.items()}


def _local(el: OxmlElement) -> str:
    return str(el.tag).rsplit("}", 1)[-1]


def fix_child_order(root: OxmlElement) -> None:
    for el in root.iter():
        order = CHILD_ORDER_INDEX.get(_local(el))
        if order is None:
            continue
        for child in sorted(el, key=lambda c: order.get(_local(c), len(order))):
            el.append(child)  # lxml: append moves an existing child to the end


def style_rpr(style: object) -> OxmlElement:
    return style.element.get_or_add_rPr()  # type: ignore[attr-defined]


def style_ppr(style: object) -> OxmlElement:
    return style.element.get_or_add_pPr()  # type: ignore[attr-defined]


def set_bidi(ppr: OxmlElement) -> None:
    _set(ppr, "w:bidi", val="1")


def setup_docx_styles(doc: Document) -> None:
    styles = doc.styles
    # Document defaults, so tables, lists and anything without a style inherit the Arabic font.
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr_default = defaults.find(qn("w:rPrDefault"))
    set_rpr(_set(rpr_default, "w:rPr"), size=DOCX_PT["body"])
    lang = _set(_set(rpr_default, "w:rPr"), "w:lang")
    lang.set(qn("w:val"), "en-GB")
    lang.set(qn("w:bidi"), "ar-SA")

    normal = styles["Normal"]
    set_rpr(style_rpr(normal), size=DOCX_PT["body"])
    set_bidi(style_ppr(normal))
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for level in (1, 2, 3):
        st = styles[f"Heading {level}"]
        rpr = style_rpr(st)
        for tag in ("w:rFonts", "w:i", "w:iCs"):
            old = rpr.find(qn(tag))
            if old is not None:
                rpr.remove(old)
        set_rpr(rpr, size=DOCX_PT[f"h{level}"], bold=True)
        st.font.color.rgb = HEADING_COLOR
        set_bidi(style_ppr(st))
        st.paragraph_format.space_before = Pt(10 if level > 1 else 0)
        st.paragraph_format.space_after = Pt(6)

    table_text = styles.add_style("Sanad Table Text", WD_STYLE_TYPE.PARAGRAPH)
    table_text.base_style = normal
    set_rpr(style_rpr(table_text), size=DOCX_PT["table"])
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.0

    notice = styles.add_style("Sanad Notice", WD_STYLE_TYPE.PARAGRAPH)
    notice.base_style = normal
    set_rpr(style_rpr(notice), size=DOCX_PT["notice"])
    ppr = style_ppr(notice)
    # In a bidi paragraph the "right" border is the start border.
    _set(_set(ppr, "w:pBdr"), "w:right", val="single", sz="18", space="6", color="B08D57")
    _set(ppr, "w:shd", val="clear", color="auto", fill="F6F1E7")

    item = styles.add_style("Sanad List Item", WD_STYLE_TYPE.PARAGRAPH)
    item.base_style = normal
    item.paragraph_format.space_after = Pt(3)


def add_runs(par: object, text: str, size: float | None = None) -> None:
    for chunk, bold in inline_runs(text):
        run = par.add_run(chunk)  # type: ignore[attr-defined]
        if bold:
            set_rpr(run._r.get_or_add_rPr(), bold=True, size=size)
        elif size is not None:
            set_rpr(run._r.get_or_add_rPr(), size=size)


def set_indent(par: object, start_mm: float, hanging_mm: float) -> None:
    ppr = par._p.get_or_add_pPr()  # type: ignore[attr-defined]
    _set(ppr, "w:ind", start=str(round(start_mm * 56.7)), hanging=str(round(hanging_mm * 56.7)))


def add_table(doc: Document, block: Block) -> None:
    table = doc.add_table(rows=1 + len(block.rows), cols=len(block.header))
    table.style = doc.styles["Table Grid"]
    tblpr = table._tbl.tblPr
    _set(tblpr, "w:bidiVisual", val="1")
    _set(tblpr, "w:tblW", w="5000", type="pct")
    for r, cells in enumerate([block.header, *block.rows]):
        for c, text in enumerate(cells):
            cell = table.cell(r, c)
            par = cell.paragraphs[0]
            par.style = doc.styles["Sanad Table Text"]
            if r == 0:
                for chunk, _ in inline_runs(text):
                    set_rpr(par.add_run(chunk)._r.get_or_add_rPr(), bold=True)
                tcpr = cell._tc.get_or_add_tcPr()
                _set(tcpr, "w:shd", val="clear", color="auto", fill="E3ECEF")
            else:
                add_runs(par, text)


def build_docx(src: Source) -> Document:
    doc = Document()
    setup_docx_styles(doc)
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)
    section.header_distance = section.footer_distance = Mm(10)
    _set(section._sectPr, "w:bidi", val="1")

    for number, text in enumerate(src.pages, 1):
        blocks = parse_blocks(text)
        first_par = None
        after_table = False
        for block in blocks:
            pars = []
            if block.kind in ("h1", "h2", "h3"):
                par = doc.add_paragraph(style=f"Heading {block.kind[1]}")
                add_runs(par, block.text)
                pars.append(par)
            elif block.kind == "p":
                par = doc.add_paragraph()
                add_runs(par, block.text)
                pars.append(par)
            elif block.kind == "quote":
                par = doc.add_paragraph(style="Sanad Notice")
                add_runs(par, block.text)
                pars.append(par)
            elif block.kind in ("ul", "ol"):
                for k, item in enumerate(block.items):
                    marker = "•" if block.kind == "ul" else f"{block.start + k}."
                    par = doc.add_paragraph(style="Sanad List Item")
                    set_indent(par, 8, 6)
                    par.add_run(f"{marker}\t")
                    add_runs(par, item)
                    pars.append(par)
            elif block.kind == "table":
                if first_par is None and number > 1:
                    raise RenderError(f"{src.doc_id} p{number}: a page may not open with a table")
                add_table(doc, block)
            if after_table and pars:
                pars[0].paragraph_format.space_before = Pt(8)
            after_table = block.kind == "table"
            if first_par is None and pars:
                first_par = pars[0]
        if number > 1:
            if first_par is None:
                raise RenderError(f"{src.doc_id} p{number}: empty page")
            first_par.paragraph_format.page_break_before = True

    fix_child_order(doc.element)
    fix_child_order(doc.styles.element)

    props = doc.core_properties
    props.title = str(src.entry["title_ar"])
    props.subject = str(src.entry["title_en"])
    props.author = props.last_modified_by = "Al-Rayhan Development Foundation (synthetic)"
    props.comments = str(load_manifest()["notice_en"])
    props.keywords = "synthetic demo corpus"
    props.created = props.modified = FIXED_DATE.replace(tzinfo=None)
    props.last_printed = FIXED_DATE.replace(tzinfo=None)
    props.revision = 1
    props.language = "ar-SA"
    props.category = str(src.entry["category"])
    return doc


def save_deterministic_zip(save: object, out: Path) -> None:
    """Save an OOXML package with fixed zip timestamps and attributes."""
    buf = io.BytesIO()
    save(buf)  # type: ignore[operator]
    buf.seek(0)
    with zipfile.ZipFile(buf) as src:
        entries = [(info.filename, src.read(info)) for info in src.infolist()]
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            dst.writestr(info, data)


def render_docx(src: Source, out: Path) -> None:
    save_deterministic_zip(build_docx(src).save, out)


# --------------------------------------------------------------------------------------------
# TXT
# --------------------------------------------------------------------------------------------


def render_txt(src: Source, out: Path) -> None:
    encoding = str(src.entry["encoding"])
    for n, page in enumerate(src.pages, 1):
        if "#" in page or "|" in page or "**" in page:
            raise RenderError(f"{src.doc_id} p{n}: Markdown markup in a TXT source")
    text = "\f".join(page.strip("\n") + "\n" for page in src.pages)
    if encoding == "windows-1256":
        # A Windows export: CRLF line ends, like Notepad wrote them.
        text = text.replace("\n", "\r\n")
        data = text.encode("cp1256")
    elif encoding == "utf-8":
        data = text.encode("utf-8")
    else:
        raise RenderError(f"{src.doc_id}: unknown encoding {encoding}")
    out.write_bytes(data)


# --------------------------------------------------------------------------------------------
# LibreOffice (as ingestion converts DOCX and TXT) and page counts
# --------------------------------------------------------------------------------------------


def decode_txt(data: bytes) -> str:
    """Ingestion's TXT fallback order (SYSTEM_DESIGN §5.2): UTF-8, UTF-8 with BOM, Windows-1256."""
    for encoding in ("utf-8", "utf-8-sig", "cp1256"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        return text.removeprefix("﻿")
    raise RenderError("TXT_ENCODING_UNKNOWN")


def libreoffice_pdf(path: Path, work: Path) -> Path:
    """Convert DOCX or TXT to PDF the way ingestion does (see seed/README.md)."""
    work.mkdir(parents=True, exist_ok=True)
    if path.suffix == ".txt":
        text = decode_txt(path.read_bytes()).replace("\r\n", "\n")
        tmp = work / path.name
        tmp.write_text(text, encoding="utf-8")
        cmd = [*LIBREOFFICE, "--infilter=Text (encoded):UTF8,LF,,,", "--convert-to", "pdf"]
        src = tmp
    else:
        cmd = [*LIBREOFFICE, "--convert-to", "pdf"]
        src = path
    subprocess.run([*cmd, "--outdir", str(work), str(src)], check=True, capture_output=True)
    out = work / (path.stem + ".pdf")
    if not out.exists():
        raise RenderError(f"LibreOffice produced no PDF for {path.name}")
    return out


def pdf_pages(path: Path) -> int:
    info = subprocess.run(["pdfinfo", str(path)], check=True, capture_output=True, text=True)
    m = re.search(r"^Pages:\s+(\d+)$", info.stdout, re.M)
    if not m:
        raise RenderError(f"pdfinfo: no page count for {path}")
    return int(m.group(1))


def pdf_fonts(path: Path) -> set[str]:
    out = subprocess.run(["pdffonts", str(path)], check=True, capture_output=True, text=True)
    names = set()
    for line in out.stdout.splitlines()[2:]:
        name = line.split()[0]
        names.add(name.split("+", 1)[-1])
    return names


# --------------------------------------------------------------------------------------------
# Previews and failure fixtures
# --------------------------------------------------------------------------------------------


def png_page(pdf: Path, page: int, out: Path, dpi: int = 100, gray: bool = False) -> None:
    stem = out.with_suffix("")
    cmd = ["pdftoppm", "-png", "-r", str(dpi), "-f", str(page), "-l", str(page), "-singlefile"]
    if gray:
        cmd.append("-gray")
    subprocess.run([*cmd, str(pdf), str(stem)], check=True)


def fake_bytes(seed: str, size: int) -> bytes:
    """Deterministic filler bytes (a sha256 chain)."""
    out = bytearray()
    block = seed.encode()
    while len(out) < size:
        block = hashlib.sha256(block).digest()
        out += block
    return bytes(out[:size])


def build_fixtures(work: Path) -> None:
    BAD_DIR.mkdir(exist_ok=True)
    base = DOCS_DIR / "annual-report-2016.pdf"

    # 1. Truncated PDF: the first 30% of a real seed PDF. The cut lands before the page tree
    #    and the cross-reference table, so no parser can rebuild the document.
    data = base.read_bytes()
    (BAD_DIR / "truncated.pdf").write_bytes(data[: len(data) * 3 // 10])

    # 2. Password-protected PDF: AES-128, a user password is needed to open it. Not AES-256,
    #    whose random file key and salts qpdf can't pin; with AES-128, a fixed ID and IV make
    #    the file reproducible. The password is a fixture value, documented in README.md.
    subprocess.run(
        [
            "qpdf", "--static-id", "--static-aes-iv",
            "--encrypt", "--user-password=FAKE-open-1234", "--owner-password=FAKE-owner-1234",
            "--bits=128", "--use-aes=y", "--", str(base), str(BAD_DIR / "encrypted.pdf"),
        ],
        check=True,
    )  # fmt: skip

    # 3. Image-only PDF, like a scan: two seed pages rasterised to greyscale PNG, no text layer.
    images = []
    for page in (1, 2):
        png = work / f"scan-{page}.png"
        png_page(base, page, png, dpi=110, gray=True)
        images.append(str(png))
    subprocess.run(
        ["img2pdf", "--nodate", "--producer=", "--pagesize=A4", "-o",
         str(BAD_DIR / "scanned-no-text.pdf"), *images],
        check=True,
    )  # fmt: skip

    # 4. Empty DOCX: a valid Word file with no text at all.
    empty = Document()
    props = empty.core_properties
    props.title = props.subject = props.comments = props.keywords = ""
    props.author = props.last_modified_by = "Al-Rayhan Development Foundation (synthetic)"
    props.created = props.modified = props.last_printed = FIXED_DATE.replace(tzinfo=None)
    props.revision = 1
    save_deterministic_zip(empty.save, BAD_DIR / "empty.docx")

    # 5. Not a PDF: bytes shaped like a Windows executable's DOS header ("MZ", e_lfanew -> "PE"),
    #    then filler. Not a runnable program, only the start of the header; named .pdf.
    header = bytearray(b"MZ" + b"\0" * 0x3A + (0x40).to_bytes(4, "little"))
    header += b"PE\0\0"
    (BAD_DIR / "not-a-pdf.pdf").write_bytes(bytes(header) + fake_bytes("sanad-not-a-pdf", 4096))


def build_previews(converted: dict[str, Path]) -> None:
    PREVIEW_DIR.mkdir(exist_ok=True)
    for file, page in PREVIEWS:
        pdf = converted.get(file, DOCS_DIR / file)
        stem = Path(file).stem
        png_page(pdf, page, PREVIEW_DIR / f"{stem}-p{page}.png")


# --------------------------------------------------------------------------------------------


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("doc_ids", nargs="*", help="render only these documents")
    parser.add_argument("--no-fixtures", action="store_true", help="skip seed/bad and previews")
    args = parser.parse_args()

    manifest = load_manifest()
    entries = manifest["documents"]  # type: ignore[index]
    if args.doc_ids:
        entries = [e for e in entries if e["doc_id"] in args.doc_ids]
        if len(entries) != len(set(args.doc_ids)):
            print("unknown doc_id in arguments", file=sys.stderr)
            return 2
    DOCS_DIR.mkdir(exist_ok=True)

    failures: list[str] = []
    converted: dict[str, Path] = {}
    with tempfile.TemporaryDirectory(prefix="sanad-render-") as tmp:
        work = Path(tmp)
        for entry in entries:
            src = load_source(entry)
            out = DOCS_DIR / str(entry["file"])
            fmt = entry["format"]
            note = ""
            if fmt == "pdf":
                shrunk = render_pdf(src, work, out)
                viewer = out
                if shrunk:
                    note = "shrunk " + ", ".join(f"p{p}->{s}pt" for p, s in sorted(shrunk.items()))
            elif fmt == "docx":
                render_docx(src, out)
                viewer = libreoffice_pdf(out, work / "lo")
            elif fmt == "txt":
                render_txt(src, out)
                viewer = libreoffice_pdf(out, work / "lo")
            else:
                raise RenderError(f"{src.doc_id}: unknown format {fmt}")
            converted[str(entry["file"])] = viewer
            pages = pdf_pages(viewer)
            fonts = sorted(pdf_fonts(viewer))
            status = "ok" if pages == entry["page_count"] else "PAGE COUNT MISMATCH"
            if status != "ok":
                failures.append(f"{entry['file']}: {pages} pages, expected {entry['page_count']}")
            if any("Kacst" in f for f in fonts):
                failures.append(f"{entry['file']}: uses KacstOne")
            print(
                f"{status:4} {entry['file']:36} {pages:>2}/{entry['page_count']:<2} "
                f"fonts={','.join(fonts)} {note}".rstrip()
            )

        if not args.no_fixtures and not args.doc_ids:
            build_fixtures(work)
            build_previews(converted)
            # Keep the scratch LibreOffice PDFs out of the repo; only previews are committed.
            shutil.rmtree(work / "lo", ignore_errors=True)

    if failures:
        print("\n".join(["FAILED:", *failures]), file=sys.stderr)
        return 1
    print(f"rendered {len(entries)} documents")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except RenderError as exc:
        print(f"render error: {exc}", file=sys.stderr)
        sys.exit(1)
